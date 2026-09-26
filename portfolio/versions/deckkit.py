# -*- coding: utf-8 -*-
"""Native PPTX toolkit — everything it makes is an editable shape or text box.

No flattened slides: fills, gradients, outlines, corner radii and copy all stay
live so the decks can be reworked in Canva or PowerPoint.
"""

import copy, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.oxml.ns import qn
from lxml import etree


# ---------------------------------------------------------------- metrics
# Design units: the slide is 1920 wide = 960pt, so 1 pt == 2 design px.
# Fonts are installed locally, so text can be measured instead of guessed.
from PIL import ImageFont as _IF

_FONT_DIR = "/root/.local/share/fonts/deck"
_FONT_FILES = {
    "Italiana": "Italiana-Regular.ttf", "Great Vibes": "GreatVibes-Regular.ttf",
    "Oswald": "Oswald.ttf", "Bodoni Moda": "BodoniModa.ttf", "Fredoka": "Fredoka.ttf",
    "VT323": "VT323-Regular.ttf", "Poppins": "Poppins-Regular.ttf",
    "Archivo Black": "ArchivoBlack-Regular.ttf", "Inter": "Inter.ttf",
    "Anton": "Anton-Regular.ttf", "Space Mono": "SpaceMono-Regular.ttf",
    "Nunito": "Nunito.ttf", "Press Start 2P": "PressStart2P-Regular.ttf",
    "Silkscreen": "Silkscreen-Regular.ttf",
}
_BOLD = {"Poppins": "Poppins-Bold.ttf", "Space Mono": "SpaceMono-Bold.ttf"}
_measure_cache = {}


def _pil(font_name, size, bold=False):
    """size is in points; the face is loaded at design-pixel scale (pt*2)."""
    fn = (_BOLD.get(font_name) if bold else None) or _FONT_FILES.get(font_name, "Poppins-Regular.ttf")
    key = (fn, int(size * 2))
    if key not in _measure_cache:
        _measure_cache[key] = _IF.truetype(os.path.join(_FONT_DIR, fn), max(4, int(size * 2)))
    return _measure_cache[key]


def measure(text, font_name="Poppins", size=14, tracking=0, bold=False):
    """Width of one line, in design px."""
    f = _pil(font_name, size, bold)
    w = f.getlength(text)
    if tracking:
        w += tracking * 2 * max(0, len(text))
    return w


def line_height(font_name="Poppins", size=14, spacing=1.4, bold=False):
    f = _pil(font_name, size, bold)
    asc, desc = f.getmetrics()
    return max(asc + desc, size * 2 * 1.2) * spacing


SAFETY = 0.92          # renderer sets slightly wider than PIL measures


def wrap(text, font_name="Poppins", size=14, width=800, tracking=0, bold=False):
    width *= SAFETY
    words, lines, cur = text.split(), [], ""
    for wd in words:
        trial = (cur + " " + wd).strip()
        if measure(trial, font_name, size, tracking, bold) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = wd
    if cur:
        lines.append(cur)
    return lines


def block_height(body, font_name="Poppins", size=14, width=800, spacing=1.4,
                 after=0, tracking=0, bold=False):
    """Predicted rendered height of a paragraph or list of paragraphs."""
    paras = body if isinstance(body, (list, tuple)) else [body]
    lh = line_height(font_name, size, spacing, bold)
    total = 0
    for i, para in enumerate(paras):
        total += len(wrap(para, font_name, size, width, tracking, bold)) * lh
        if i < len(paras) - 1:
            total += after * 2
    return total


def fit_size(text, font_name="Poppins", max_w=800, max_h=None, start=40, min_size=8,
             spacing=1.25, tracking=0, bold=False, step=0.5):
    """Largest size at which the text fits the box."""
    size = start
    while size > min_size:
        paras = text if isinstance(text, (list, tuple)) else [text]
        ok = True
        for para in paras:
            if measure(para, font_name, size, tracking, bold) > max_w and " " not in para:
                ok = False
        if ok:
            h = block_height(text, font_name, size, max_w, spacing, 0, tracking, bold)
            if max_h is None or h <= max_h:
                return size
        size -= step
    return min_size

W, H = 1920, 1080

def px(v):
    return Emu(int(round(v * 6350)))

def C(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def _shadow(shape, blur=28, dist=10, direction=5400000, color="000000", alpha=22):
    """Outer shadow via DrawingML (PowerPoint keeps it live and editable)."""
    spPr = shape._element.spPr
    for tag in ("a:effectLst",):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    eff = etree.SubElement(spPr, "{%s}effectLst" % ns)
    sh = etree.SubElement(eff, "{%s}outerShdw" % ns)
    sh.set("blurRad", str(int(blur * 12700)))
    sh.set("dist", str(int(dist * 12700)))
    sh.set("dir", str(int(direction)))
    sh.set("rotWithShape", "0")
    clr = etree.SubElement(sh, "{%s}srgbClr" % ns)
    clr.set("val", color.lstrip("#").upper())
    a = etree.SubElement(clr, "{%s}alpha" % ns)
    a.set("val", str(int(alpha * 1000)))


class Slide(object):
    def __init__(self, sl):
        self.sl = sl

    # ---------------------------------------------------------------- shapes
    def shape(self, kind, x, y, w, h, fill=None, grad=None, angle=90, line=None,
              lw=1.5, radius=None, rot=0, dash=None, shadow=None, adj=None, alpha=None):
        s = self.sl.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
        s.shadow.inherit = False
        if radius is not None:
            try:
                s.adjustments[0] = radius
            except Exception:
                pass
        if adj is not None:
            for i, v in enumerate(adj):
                try:
                    s.adjustments[i] = v
                except Exception:
                    pass
        if grad:
            s.fill.gradient()
            stops = s.fill.gradient_stops
            stops[0].color.rgb = C(grad[0])
            stops[0].position = 0.0
            stops[1].color.rgb = C(grad[1])
            stops[1].position = 1.0
            s.fill.gradient_angle = angle
        elif fill:
            s.fill.solid()
            s.fill.fore_color.rgb = C(fill)
            if alpha is not None:
                ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
                clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                a = etree.SubElement(clr, "{%s}alpha" % ns)
                a.set("val", str(int(alpha * 1000)))
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = C(line)
            s.line.width = Pt(lw)
            if dash:
                s.line.dash_style = dash
        else:
            s.line.fill.background()
        if rot:
            s.rotation = rot
        if shadow:
            _shadow(s, **shadow) if isinstance(shadow, dict) else _shadow(s)
        s.text_frame.word_wrap = True
        return s

    def rect(self, *a, **k):
        return self.shape(MSO_SHAPE.RECTANGLE, *a, **k)

    def rrect(self, *a, **k):
        k.setdefault("radius", 0.08)
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, *a, **k)

    def oval(self, *a, **k):
        return self.shape(MSO_SHAPE.OVAL, *a, **k)

    def bg(self, fill=None, grad=None, angle=90):
        return self.rect(0, 0, W, H, fill=fill, grad=grad, angle=angle)

    def line(self, x, y, length, color="000000", weight=1.2, vertical=False, dash=None):
        s = self.rect(x, y, weight if vertical else length, length if vertical else weight,
                      fill=color)
        if dash:
            s.line.color.rgb = C(color)
            s.line.width = Pt(weight)
            s.line.dash_style = dash
        return s

    # ---------------------------------------------------------------- text
    def text(self, x, y, w, h, body, font="Poppins", size=14, bold=False, italic=False,
             color="000000", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.4,
             tracking=0, after=8, shape=None, caps=False, wrap_text=True):
        if shape is None:
            tb = self.sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
        else:
            tb = shape
        tf = tb.text_frame
        tf.word_wrap = wrap_text
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        lines = body if isinstance(body, (list, tuple)) else [body]
        for i, t in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            p.space_after = Pt(after if len(lines) > 1 else 0)
            r = p.add_run()
            r.text = t.upper() if caps else t
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = C(color)
            if tracking:
                r.font._rPr.set("spc", str(int(tracking * 100)))
        return tb

    def label(self, x, y, w, text, size=12, color="000000", tracking=6, font="Oswald",
              align=PP_ALIGN.LEFT, bold=False, caps=True):
        return self.text(x, y, w, size * 2.4, text, font=font, size=size, bold=bold,
                         color=color, align=align, tracking=tracking, caps=caps)

    # ---------------------------------------------------------------- parts
    def photo(self, x, y, w, h, label="drop photo", fill="EDEDF2", line="B9B4C6",
              radius=None, rot=0, dash=True, font="Poppins", size=10,
              label_color="8C8698", shadow=None):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
        s = self.shape(kind, x, y, w, h, fill=fill, line=line, lw=1.25, radius=radius,
                       rot=rot, dash=MSO_LINE_DASH_STYLE.DASH if dash else None,
                       shadow=shadow)
        p = s.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = p.add_run()
        r.text = label
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = C(label_color)
        return s

    def pill(self, x, y, text, size=12, pad=34, h=44, fill="FFFFFF", color="000000",
             line=None, font="Oswald", tracking=5, bold=False, radius=0.5, width=None):
        w = width or (pad * 2 + measure(text, font, size, tracking, bold) / SAFETY)
        s = self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line,
                       lw=1.25, radius=radius)
        tf = s.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = C(color)
        if tracking:
            r.font._rPr.set("spc", str(int(tracking * 100)))
        return w

    def window(self, x, y, w, h, title, bar="FF6FB5", bar2=None, body="FFFFFF",
               radius=0.05, bar_h=46, title_font="Poppins", title_size=13,
               title_color="FFFFFF", dots=True, dot_colors=("FFFFFF",) * 3,
               line=None, shadow=True, tracking=2):
        """A window card: frame, title bar, buttons — all separate editable shapes."""
        frame = self.rrect(x, y, w, h, fill=body, line=line, lw=1.5, radius=radius,
                           shadow={} if shadow else None)
        if bar2:
            self.rrect(x, y, w, bar_h * 2, grad=(bar, bar2), angle=90, radius=radius * 2)
        else:
            self.rrect(x, y, w, bar_h * 2, fill=bar, radius=radius * 2)
        self.rect(x, y + bar_h - 2, w, 2, fill=bar2 or bar)
        self.text(x + 22, y, w - 140, bar_h, title, font=title_font, size=title_size,
                  color=title_color, anchor=MSO_ANCHOR.MIDDLE, tracking=tracking)
        if dots:
            for i, c in enumerate(dot_colors):
                self.oval(x + w - 34 - i * 30, y + bar_h / 2 - 9, 18, 18, fill=c)
        return frame, (x, y + bar_h, w, h - bar_h)

    def bubble(self, x, y, w, h, text, fill="FFFFFF", line="000000", lw=2.5, color="000000",
               font="Poppins", size=14, tail="bl", radius=0.28, spacing=1.45, pad=34):
        self.rrect(x, y, w, h, fill=fill, line=line, lw=lw, radius=radius)
        if tail:
            tx = x + 44 if tail.endswith("l") else x + w - 76
            t = self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE, tx, y + h - 4, 34, 26,
                           fill=fill, line=line, lw=lw)
            t.rotation = 180
            self.rect(tx + 2, y + h - 8, 30, 8, fill=fill)
        self.text(x + pad, y + 18, w - pad * 2, h - 34, text, font=font, size=size,
                  color=color, spacing=spacing, anchor=MSO_ANCHOR.MIDDLE)



    def head(self, x, y, text, font="Italiana", size=40, color="000000", tracking=0,
             caps=False, max_w=None, spacing=1.15, align=PP_ALIGN.LEFT, bold=False,
             italic=False, gap=0, one_line=False, min_size=10):
        """Draw a heading, auto-shrunk to fit max_w, and return its bottom edge."""
        txt = text.upper() if caps else text
        if max_w and one_line:
            while size > min_size and measure(txt, font, size, tracking, bold) > max_w * SAFETY:
                size -= 0.5
        if one_line:
            lines, box_w = [txt], max(max_w or 0, measure(txt, font, size, tracking, bold) + 80)
        else:
            lines = wrap(txt, font, size, max_w or 4000, tracking, bold)
            box_w = max_w or 1700
        lh = line_height(font, size, spacing, bold)
        h = lh * len(lines) + lh * 0.25
        self.text(x, y, box_w, h, txt, font=font, size=size, color=color,
                  tracking=tracking, spacing=spacing, align=align, bold=bold,
                  italic=italic, wrap_text=not one_line)
        return y + h + gap

    def para(self, x, y, w, body, font="Poppins", size=13, color="000000", spacing=1.7,
             after=16, max_h=None, italic=False, align=PP_ALIGN.LEFT, tracking=0,
             gap=0, min_size=8):
        """Draw body copy, shrinking only if it would not fit max_h. Returns bottom."""
        if max_h:
            while size > min_size and block_height(body, font, size, w, spacing, after,
                                                   tracking) > max_h:
                size -= 0.5
        h = block_height(body, font, size, w, spacing, after, tracking)
        self.text(x, y, w, h + 10, body, font=font, size=size, color=color,
                  spacing=spacing, after=after, italic=italic, align=align,
                  tracking=tracking)
        return y + h + gap

    # ------------------------------------------------------------ imagery
    SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rrect": MSO_SHAPE.ROUNDED_RECTANGLE,
              "oval": MSO_SHAPE.OVAL, "hex": MSO_SHAPE.HEXAGON,
              "oct": MSO_SHAPE.OCTAGON, "para": MSO_SHAPE.PARALLELOGRAM,
              "trap": MSO_SHAPE.TRAPEZOID, "diamond": MSO_SHAPE.DIAMOND,
              "pent": MSO_SHAPE.PENTAGON, "plaque": MSO_SHAPE.PLAQUE,
              "arch": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
              "snip": MSO_SHAPE.SNIP_2_DIAG_RECTANGLE,
              "snip1": MSO_SHAPE.SNIP_1_RECTANGLE,
              "round1": MSO_SHAPE.ROUND_1_RECTANGLE,
              "tag": MSO_SHAPE.PENTAGON, "heart": MSO_SHAPE.HEART,
              "cloud": MSO_SHAPE.CLOUD, "star": MSO_SHAPE.STAR_5_POINT}

    def cut(self, kind, x, y, w, h, label="drop photo", fill="E4E6EE", line=None,
            lw=3, rot=0, radius=None, label_color="7C8398", size=10, shadow=None,
            font="Poppins", dash=False):
        """A picture placeholder that is NOT a plain rectangle — picture-fill it
        in Canva/PowerPoint and the silhouette is kept."""
        s = self.shape(self.SHAPES.get(kind, MSO_SHAPE.RECTANGLE), x, y, w, h,
                       fill=fill, line=line, lw=lw, radius=radius, rot=rot,
                       shadow=shadow,
                       dash=MSO_LINE_DASH_STYLE.DASH if dash else None)
        tf = s.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = C(label_color)
        return s

    def annot(self, x, y, text, color="000000", size=10, font="Space Mono",
              leader=0, tracking=3, align=PP_ALIGN.LEFT, width=320):
        """A small annotation, optionally with a leader rule."""
        if leader:
            self.rect(x, y + size * 1.5, leader, 1.5, fill=color)
        return self.text(x + (leader + 10 if leader else 0), y, width, size * 2.6, text,
                         font=font, size=size, color=color, tracking=tracking,
                         align=align)

    def tape(self, x, y, w=150, h=40, rot=-8, fill="F4E6A8", alpha=62):
        return self.rect(x, y, w, h, fill=fill, rot=rot, alpha=alpha)

    def pic(self, path, x, y, w, h=None):
        if h is None:
            from PIL import Image
            im = Image.open(path)
            h = w * im.size[1] / im.size[0]
        return self.sl.shapes.add_picture(path, px(x), px(y), px(w), px(h))

    def star(self, x, y, size, fill="FFE83D", line=None, lw=2, points=5, rot=0):
        kind = {4: MSO_SHAPE.STAR_4_POINT, 5: MSO_SHAPE.STAR_5_POINT,
                8: MSO_SHAPE.STAR_8_POINT, 16: MSO_SHAPE.STAR_16_POINT,
                24: MSO_SHAPE.STAR_24_POINT}[points]
        return self.shape(kind, x, y, size, size, fill=fill, line=line, lw=lw, rot=rot)

    def heart(self, x, y, size, fill="FF2E93", line=None, lw=2, rot=0):
        return self.shape(MSO_SHAPE.HEART, x, y, size, size * 0.92, fill=fill, line=line,
                          lw=lw, rot=rot)


class Deck(object):
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = px(W), px(H)
        self.blank = self.prs.slide_layouts[6]

    def new(self):
        return Slide(self.prs.slides.add_slide(self.blank))

    def save(self, path):
        self.prs.save(path)
        return path
