# -*- coding: utf-8 -*-
"""Zahra Levina — 2026 portfolio deck.

A 16:9 (1920x1080) presentation built as a 2000s desktop OS: every slide is a
window on a patterned desktop, with a persistent taskbar binding the two
palettes (candy pink / frutiger aero) into one machine.

Layout rule: the slide is 13.333in wide, so 1pt of type == 2 design px.
All vertical rhythm below is derived from that.
"""

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

def px(v):
    return Emu(int(round(v * 6350)))

def C(h):
    return RGBColor.from_string(h)

PINK = {"desk_a": "FFE3F1", "desk_b": "F7D2FF", "bar_a": "FF6FB5", "bar_b": "C77DFF",
        "accent": "B49CFF", "pop": "C6F24E", "soft": "FFF1F8", "edge": "E36BA6",
        "chip": "FFDCEE"}
AERO = {"desk_a": "E4F6BE", "desk_b": "9BD4F5", "bar_a": "7FD4F7", "bar_b": "3A8FD6",
        "accent": "5FB8E8", "pop": "FF6FB5", "soft": "EFF8FF", "edge": "2F7CBF",
        "chip": "DCEFFB"}
INK, MUTE = "2B2340", "6E6486"
HEAD, UI, MONO, BODY = "Arial Black", "Tahoma", "Courier New", "Verdana"

TASKS = ["portfolio_2026.exe", "about_zahra.txt", "control_panel", "experience",
         "sawalaspace.com", "my_creative_works", "stretch_for_stray.html",
         "beyond_creative", "skills_and_tools", "new_message"]

# ------------------------------------------------------------- primitives
def shape(sl, kind, x, y, w, h):
    s = sl.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    s.shadow.inherit = False
    return s

def solid(s, c):
    s.fill.solid(); s.fill.fore_color.rgb = C(c)

def grad(s, c1, c2, angle=90):
    s.fill.gradient()
    st = s.fill.gradient_stops
    st[0].color.rgb = C(c1); st[0].position = 0.0
    st[1].color.rgb = C(c2); st[1].position = 1.0
    s.fill.gradient_angle = angle

def noline(s):
    s.line.fill.background()

def stroke(s, c, w=1.5):
    s.line.color.rgb = C(c); s.line.width = Pt(w)

def radius(s, r):
    try:
        s.adjustments[0] = r
    except Exception:
        pass

def deco(sl, kind, x, y, w, h, fill, rot=0, edge=None):
    """A drawn ornament — never a font glyph, so it renders everywhere."""
    s = shape(sl, kind, x, y, w, h)
    solid(s, fill)
    if edge:
        stroke(s, edge, 1.25)
    else:
        noline(s)
    if rot:
        s.rotation = rot
    return s

def label(sl, x, y, w, h, text, font=BODY, size=18, bold=False, color=INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, spc=0):
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = C(color)
    if spc:
        r.font._rPr.set("spc", str(int(spc * 100)))
    return tf

def paras(sl, x, y, w, h, items, font=BODY, size=17, color=INK, spacing=1.55, gap=18):
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = t
        r.font.name = font; r.font.size = Pt(size); r.font.color.rgb = C(color)
    return tf

# ------------------------------------------------------------- chrome
def desktop(sl, pal):
    bg = shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 1920, 1080)
    grad(bg, pal["desk_a"], pal["desk_b"], 45); noline(bg)
    for sx, sy, sz in [(72, 150, 30), (1826, 210, 24), (236, 878, 26), (1770, 842, 32),
                       (44, 548, 20), (1870, 560, 18), (436, 74, 20), (1508, 92, 22),
                       (972, 32, 16), (694, 962, 20), (1298, 972, 18)]:
        deco(sl, MSO_SHAPE.STAR_4_POINT, sx, sy, sz, sz, pal["bar_a"])

def taskbar(sl, pal, idx):
    bar = shape(sl, MSO_SHAPE.RECTANGLE, 0, 1016, 1920, 64)
    grad(bar, pal["bar_a"], pal["bar_b"], 90); noline(bar)
    deco(sl, MSO_SHAPE.RECTANGLE, 0, 1016, 1920, 3, "FFFFFF")

    start = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 18, 1026, 150, 44)
    radius(start, 0.35); grad(start, "FFFFFF", pal["chip"], 90); stroke(start, "FFFFFF", 1.25)
    tri = deco(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, 38, 1039, 18, 18, pal["bar_a"])
    tri.rotation = 90
    label(sl, 66, 1026, 90, 44, "start", UI, 16, True, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

    chip = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 186, 1026, 400, 44)
    radius(chip, 0.3); solid(chip, pal["soft"]); stroke(chip, "FFFFFF", 1.25)
    label(sl, 196, 1026, 380, 44, TASKS[idx - 1], UI, 14, False, INK,
          PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    deco(sl, MSO_SHAPE.HEART, 1516, 1038, 24, 22, "FFFFFF")
    deco(sl, MSO_SHAPE.STAR_5_POINT, 1554, 1036, 26, 26, "FFFFFF")
    deco(sl, MSO_SHAPE.STAR_4_POINT, 1594, 1038, 22, 22, "FFFFFF")

    clock = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 1602, 1026, 300, 44)
    radius(clock, 0.3); solid(clock, pal["soft"]); stroke(clock, "FFFFFF", 1.25)
    label(sl, 1612, 1026, 280, 44, "%02d/10  ·  10:30 PM" % idx, UI, 14, False, INK,
          PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def window(sl, x, y, w, h, title, pal, bar_h=56, body="FFFFFF"):
    frame = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    radius(frame, 0.035); solid(frame, body); stroke(frame, pal["edge"], 2)
    tb = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, bar_h)
    radius(tb, 0.3); grad(tb, pal["bar_a"], pal["bar_b"], 90); noline(tb)
    deco(sl, MSO_SHAPE.RECTANGLE, x, y + bar_h - 18, w, 18, pal["bar_b"])
    deco(sl, MSO_SHAPE.RECTANGLE, x, y + bar_h, w, 2, pal["edge"])
    label(sl, x + 24, y, w - 220, bar_h, title, UI, 16, True, "FFFFFF",
          PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    bx = x + w - 132
    for i in range(3):
        b = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx + i * 38, y + 14, 28, 28)
        radius(b, 0.25); solid(b, "FFFFFF"); stroke(b, pal["edge"], 1)
    deco(sl, MSO_SHAPE.RECTANGLE, bx + 7, y + 34, 14, 3, INK)                 # _
    m = deco(sl, MSO_SHAPE.RECTANGLE, bx + 45, y + 21, 14, 14, "FFFFFF")      # square
    stroke(m, INK, 1)
    deco(sl, MSO_SHAPE.RECTANGLE, bx + 82, y + 26, 16, 3, INK, 45)            # x
    deco(sl, MSO_SHAPE.RECTANGLE, bx + 82, y + 26, 16, 3, INK, -45)
    return x, y + bar_h + 2, w, h - bar_h - 2

def ph_frame(sl, x, y, w, h, caption, pal, rot=0, tint=None, cap_size=12.5):
    f = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    radius(f, 0.05); solid(f, tint or pal["soft"]); stroke(f, pal["accent"], 1.75)
    f.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if rot:
        f.rotation = rot
    star = deco(sl, MSO_SHAPE.STAR_4_POINT, x + w / 2 - 17, y + h / 2 - 64, 34, 34, pal["accent"])
    if rot:
        star.rotation = rot
    label(sl, x + 16, y + h / 2 - 18, w - 32, 70, caption, UI, cap_size, True, MUTE,
          PP_ALIGN.CENTER, MSO_ANCHOR.TOP, spacing=1.35)
    return f

def sticker(sl, x, y, size, kind, color, rot=0):
    s = shape(sl, MSO_SHAPE.OVAL, x, y, size, size)
    solid(s, "FFFFFF"); stroke(s, color, 1.75)
    inner = size * 0.5
    deco(sl, kind, x + (size - inner) / 2, y + (size - inner) / 2, inner, inner, color, rot)

def chips(sl, x, y, items, pal, size=15, gap=12, h=42, pad=26):
    cx = x
    for it in items:
        w = pad * 2 + len(it) * size * 1.12
        c = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, y, w, h)
        radius(c, 0.5); solid(c, pal["chip"]); stroke(c, pal["accent"], 1.25)
        label(sl, cx, y, w, h, it, UI, size, True, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        cx += w + gap
    return cx

def holo(sl, x, y, w, h, text):
    b = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    radius(b, 0.5); grad(b, "FFB3DE", "9BD4F5", 0); noline(b)
    label(sl, x, y, w, h, text, UI, 13, True, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def checker(sl, x, y, w, cell=16):
    for i in range(int(w // cell)):
        deco(sl, MSO_SHAPE.RECTANGLE, x + i * cell, y, cell, cell,
             INK if i % 2 == 0 else "FFFFFF")

def eyebrow(sl, x, y, w, text, pal, size=15):
    label(sl, x, y, w, 30, text, MONO, size, True, pal["edge"], PP_ALIGN.LEFT,
          MSO_ANCHOR.MIDDLE, spc=1.2)

def rule(sl, x, y, w, color):
    deco(sl, MSO_SHAPE.RECTANGLE, x, y, w, 5, color)

# ------------------------------------------------------------- deck
prs = Presentation()
prs.slide_width, prs.slide_height = px(1920), px(1080)
BLANK = prs.slide_layouts[6]

def new(pal):
    sl = prs.slides.add_slide(BLANK)
    desktop(sl, pal)
    return sl

# --- 1 · cover -----------------------------------------------------------
sl = new(PINK)
checker(sl, 0, 0, 1920)
bx, by, bw, bh = window(sl, 96, 132, 1160, 792, "portfolio_2026.exe", PINK)
eyebrow(sl, bx + 64, by + 34, 600, "C:\\> 2026 PORTFOLIO", PINK)
label(sl, bx + 60, by + 80, 1040, 190, "ZAHRA", HEAD, 88, True, INK, spacing=1.0)
label(sl, bx + 60, by + 250, 1040, 190, "LEVINA", HEAD, 88, True, PINK["bar_a"], spacing=1.0)
rule(sl, bx + 64, by + 452, 420, PINK["pop"])
label(sl, bx + 62, by + 482, 1040, 90,
      "creative · content · marketing · visual communication", UI, 16, True, INK, spacing=1.4)
chips(sl, bx + 62, by + 578, ["WhatsApp", "Email", "LinkedIn"], PINK, 16, 16, 48)
label(sl, bx + 64, by + 640, 1040, 30,
      "+62 8xx-xxxx-xxxx  ·  hello@email.com  ·  linkedin.com/in/username", MONO, 13, False, MUTE)
holo(sl, bx + 64, by + 692, 1032, 34, "welcome 2 my portfolio  ·  content · design · events  ·  est. 2026")

pol = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 1348, 214, 470, 580)
radius(pol, 0.04); solid(pol, "FFFFFF"); stroke(pol, PINK["edge"], 2); pol.rotation = -4
ph_frame(sl, 1378, 250, 410, 412, "drop your portrait here\n(410 × 412)", PINK, rot=-4)
label(sl, 1378, 690, 410, 50, "me.jpg", UI, 19, True, INK, PP_ALIGN.CENTER)
sticker(sl, 1294, 176, 92, MSO_SHAPE.HEART, PINK["bar_a"], -12)
sticker(sl, 1792, 704, 78, MSO_SHAPE.STAR_5_POINT, PINK["accent"], 10)
sticker(sl, 1248, 762, 70, MSO_SHAPE.SUN, PINK["pop"])
taskbar(sl, PINK, 1)

# --- 2 · about -----------------------------------------------------------
sl = new(PINK)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "about_zahra.txt  —  Notepad", PINK)
eyebrow(sl, bx + 68, by + 34, 520, "file · edit · format · view", PINK, 14)
label(sl, bx + 64, by + 78, 1180, 130, "HI, I'M ZAHRA.", HEAD, 52, True, INK, spacing=1.0)
rule(sl, bx + 68, by + 216, 300, PINK["bar_a"])
paras(sl, bx + 66, by + 254, 1160, 520, [
    "I'm a creative and marketing enthusiast with experience in content creation, social media, graphic design, event management, and brand partnerships.",
    "With a background in agricultural community development, I've had the opportunity to work across different environments—from student organizations and research projects to building a wellness community through Sawala Space.",
    "I enjoy turning ideas into clear, engaging, and purposeful creative work.",
], size=15.5, spacing=1.4, gap=18)

cx, cy, cw, ch = bx + 1252, by + 66, 420, 530
cam = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
radius(cam, 0.05); solid(cam, PINK["soft"]); stroke(cam, PINK["edge"], 2)
deco(sl, MSO_SHAPE.RECTANGLE, cx, cy, cw, 40, PINK["bar_a"])
label(sl, cx + 16, cy, cw - 32, 40, "webcam.exe", UI, 14, True, "FFFFFF", PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
deco(sl, MSO_SHAPE.OVAL, cx + cw - 42, cy + 15, 12, 12, "FFFFFF")
ph_frame(sl, cx + 26, cy + 64, cw - 52, 380, "drop a photo here\n(368 × 380)", PINK)
label(sl, cx + 26, cy + 462, cw - 52, 40, "say hi", UI, 15, True, MUTE, PP_ALIGN.CENTER)
sticker(sl, cx - 44, cy + ch - 66, 76, MSO_SHAPE.STAR_5_POINT, PINK["accent"], -10)
taskbar(sl, PINK, 2)

# --- 3 · what I do -------------------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "control_panel  —  what i do", AERO)
eyebrow(sl, bx + 68, by + 32, 700, "select a category", AERO)
label(sl, bx + 64, by + 72, 1200, 110, "WHAT I DO", HEAD, 56, True, INK, spacing=1.0)
cards = [("Content & Social Media", "Content creation · social media · copywriting", MSO_SHAPE.STAR_4_POINT),
         ("Creative & Visual", "Graphic design · visual communication · campaign materials", MSO_SHAPE.DIAMOND),
         ("Marketing & Events", "Campaign development · event management · brand partnerships", MSO_SHAPE.STAR_5_POINT),
         ("Communication", "Community engagement · collaboration · project coordination", MSO_SHAPE.HEART)]
cwd, chd, gx, gy = 792, 262, 44, 32
for i, (title, sub, icon) in enumerate(cards):
    cx = bx + 66 + (i % 2) * (cwd + gx)
    cy = by + 198 + (i // 2) * (chd + gy)
    card = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cwd, chd)
    radius(card, 0.08); solid(card, AERO["soft"]); stroke(card, AERO["accent"], 1.75)
    ic = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx + 36, cy + 48, 88, 88)
    radius(ic, 0.22); grad(ic, AERO["bar_a"], AERO["bar_b"], 90); noline(ic)
    deco(sl, icon, cx + 58, cy + 70, 44, 44, "FFFFFF")
    label(sl, cx + 152, cy + 48, cwd - 200, 100, title, HEAD, 22, True, INK, spacing=1.15)
    label(sl, cx + 154, cy + 152, cwd - 204, 90, sub, BODY, 15.5, False, MUTE, spacing=1.45)
    label(sl, cx + cwd - 86, cy + 18, 60, 30, "0%d" % (i + 1), MONO, 14, True, AERO["accent"], PP_ALIGN.RIGHT)
sticker(sl, 1788, 148, 84, MSO_SHAPE.STAR_4_POINT, AERO["pop"])
taskbar(sl, AERO, 3)

# --- 4 · experience ------------------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "C:\\ zahra \\ experience", AERO)
deco(sl, MSO_SHAPE.RECTANGLE, bx, by, 400, bh, AERO["chip"])
deco(sl, MSO_SHAPE.RECTANGLE, bx + 400, by, 2, bh, AERO["accent"])
eyebrow(sl, bx + 34, by + 36, 340, "quick links", AERO, 14)
for i, item in enumerate(["all folders", "by year", "by role", "creative work", "organizations"]):
    iy = by + 86 + i * 52
    deco(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, bx + 38, iy + 15, 14, 14, AERO["edge"], 90)
    label(sl, bx + 64, iy, 320, 44, item, UI, 15, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
note = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx + 30, by + 410, 340, 210)
radius(note, 0.08); solid(note, "FFFDF0"); stroke(note, AERO["accent"], 1.25)
label(sl, bx + 56, by + 440, 290, 160, "note to self\n\n5 folders\nlast modified 2026",
      UI, 14, False, MUTE, spacing=1.5)
deco(sl, MSO_SHAPE.STAR_4_POINT, bx + 320, by + 430, 26, 26, AERO["pop"])

label(sl, bx + 452, by + 36, 1200, 110, "WHERE I'VE BEEN", HEAD, 44, True, INK, spacing=1.0)
rows = [("Sawala Space", "Co-Founder · Creative & Marketing"),
        ("Kementerian Pertanian RI", "Intern · Administration & Partnership Support"),
        ("IPB University", "Research Assistant"),
        ("Nabila Farm Lembang", "Intern · Agriculture & Content Creation"),
        ("Himpunan Mahasiswa PPP", "Head of Media & Branding")]
for i, (org, role) in enumerate(rows):
    ry = by + 152 + i * 120
    r = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx + 442, ry, 1244, 102)
    radius(r, 0.14); solid(r, "FFFFFF" if i % 2 else AERO["soft"]); stroke(r, AERO["accent"], 1.25)
    fold = shape(sl, MSO_SHAPE.FOLDED_CORNER, bx + 474, ry + 24, 72, 56)
    grad(fold, AERO["bar_a"], AERO["bar_b"], 90); noline(fold)
    label(sl, bx + 576, ry + 14, 980, 48, org, HEAD, 23, True, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    label(sl, bx + 578, ry + 60, 980, 32, role, BODY, 15.5, False, MUTE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    deco(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, bx + 1610, ry + 42, 18, 18, AERO["accent"], 90)
taskbar(sl, AERO, 4)

# --- 5 · Sawala Space ----------------------------------------------------
sl = new(PINK)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "sawalaspace.com", PINK)
url = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx + 30, by + 20, 1668, 50)
radius(url, 0.5); solid(url, PINK["soft"]); stroke(url, PINK["accent"], 1.25)
label(sl, bx + 62, by + 20, 1600, 50, "https://www.sawalaspace.com", UI, 15, False, MUTE,
      PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
label(sl, bx + 64, by + 96, 1100, 130, "SAWALA SPACE", HEAD, 54, True, PINK["bar_a"], spacing=1.0)
label(sl, bx + 66, by + 236, 1100, 50, "Wellness Community & Event Organizer", UI, 23, True, INK)
rule(sl, bx + 68, by + 310, 280, PINK["pop"])
label(sl, bx + 66, by + 336, 1040, 240,
      "As a co-founder, I contribute to the creative and marketing side of Sawala Space, from developing event concepts and promotional content to managing brand partnerships and supporting event execution.",
      BODY, 16, False, INK, spacing=1.55)
eyebrow(sl, bx + 68, by + 600, 400, "MY ROLE", PINK, 14)
chips(sl, bx + 66, by + 638, ["Content", "Social Media", "Campaign"], PINK, 15, 12, 44)
chips(sl, bx + 66, by + 692, ["Event", "Partnership", "Branding"], PINK, 15, 12, 44)
ph_frame(sl, bx + 1180, by + 104, 470, 300, "Sawala Space visual\n(470 × 300)", PINK)
ph_frame(sl, bx + 1180, by + 424, 225, 236, "event photo", PINK)
ph_frame(sl, bx + 1425, by + 424, 225, 236, "community", PINK)
sticker(sl, bx + 1126, by + 66, 80, MSO_SHAPE.HEART, PINK["bar_a"], -10)
taskbar(sl, PINK, 5)

# --- 6 · gallery ---------------------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "my_creative_works  —  Gallery", AERO)
label(sl, bx + 64, by + 32, 1520, 90, "SELECTED CREATIVE WORKS", HEAD, 40, True, INK, spacing=1.0)
label(sl, bx + 66, by + 118, 1560, 80,
      "A selection of social media content, event materials, promotional visuals, and other creative work I've developed across different projects.",
      BODY, 15, False, MUTE, spacing=1.5)
deco(sl, MSO_SHAPE.RECTANGLE, bx, by + 240, bw, 66, AERO["chip"])
chips(sl, bx + 62, by + 252, ["Social Media", "Event Visuals", "Promotional Materials", "Branding"],
      AERO, 14, 10, 42)
for k in range(3):
    m = deco(sl, MSO_SHAPE.RECTANGLE, bx + bw - 150 + k * 34, by + 262, 22, 22, "FFFFFF")
    stroke(m, AERO["edge"], 1.25)
caps = ["IG carousel", "event poster", "reels cover", "brand kit",
        "feed layout", "promo flyer", "campaign visual", "merch design"]
tw, th = 386, 164
for i, cap in enumerate(caps):
    cx = bx + 62 + (i % 4) * (tw + 24)
    cy = by + 326 + (i // 4) * (th + 64)
    ph_frame(sl, cx, cy, tw, th, cap, AERO, cap_size=12)
    label(sl, cx, cy + th + 8, tw, 28, cap.replace(" ", "_") + ".jpg", MONO, 12.5, False, MUTE, PP_ALIGN.CENTER)
taskbar(sl, AERO, 6)

# --- 7 · Stretch for Stray -----------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "stretch_for_stray.html", AERO)
ph_frame(sl, bx + 54, by + 40, 470, 620, "event poster\n(470 × 620)", AERO)
label(sl, bx + 54, by + 672, 470, 36, "poster_final_v3.png", MONO, 12.5, False, MUTE, PP_ALIGN.CENTER)
tx = bx + 570
eyebrow(sl, tx + 4, by + 40, 600, "Wellness × Social Impact", AERO, 14)
label(sl, tx, by + 74, 1150, 110, "STRETCH FOR STRAY", HEAD, 42, True, INK, spacing=1.0)
rule(sl, tx + 4, by + 192, 260, AERO["pop"])
label(sl, tx, by + 226, 1100, 110,
      "A wellness event combining movement, community, and support for animal welfare.",
      BODY, 17.5, False, INK, spacing=1.5)
eyebrow(sl, tx + 4, by + 378, 400, "MY CONTRIBUTION", AERO, 14)
for i, it in enumerate(["Campaign concept", "Event promotion", "Social media content",
                        "Partnership communication", "Event coordination"]):
    iy = by + 420 + i * 56
    b = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, tx + 2, iy, 34, 34)
    radius(b, 0.3); grad(b, AERO["bar_a"], AERO["bar_b"], 90); noline(b)
    label(sl, tx + 2, iy, 34, 34, "0%d" % (i + 1), MONO, 12, True, "FFFFFF",
          PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    label(sl, tx + 52, iy - 2, 600, 40, it, BODY, 17.5, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

mx, my = tx + 700, by + 366
mp = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, mx, my, 390, 190)
radius(mp, 0.1); solid(mp, AERO["soft"]); stroke(mp, AERO["accent"], 1.5)
deco(sl, MSO_SHAPE.RECTANGLE, mx, my, 390, 36, AERO["bar_b"])
label(sl, mx + 16, my, 360, 36, "now playing", UI, 13.5, True, "FFFFFF", PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
label(sl, mx + 26, my + 58, 340, 32, "stretch_for_stray.mp4", UI, 13, True, INK)
label(sl, mx + 26, my + 92, 340, 26, "documentation reel", BODY, 12.5, False, MUTE)
trk = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, mx + 26, my + 134, 336, 10)
radius(trk, 0.5); solid(trk, "FFFFFF"); stroke(trk, AERO["accent"], 1)
fb = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, mx + 26, my + 134, 210, 10)
radius(fb, 0.5); grad(fb, AERO["bar_a"], AERO["pop"], 0); noline(fb)
for k, ox in enumerate([148, 186, 222]):
    t = deco(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, mx + ox, my + 156, 20, 20, AERO["edge"])
    t.rotation = 270 if k == 0 else 90
ph_frame(sl, mx, by + 578, 186, 130, "event docs", AERO, cap_size=11)
ph_frame(sl, mx + 204, by + 578, 186, 130, "social post", AERO, cap_size=10)
sticker(sl, bx + 482, by + 16, 76, MSO_SHAPE.STAR_5_POINT, AERO["pop"], -12)
taskbar(sl, AERO, 7)

# --- 8 · beyond creative work --------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "beyond_creative_work", AERO)
label(sl, bx + 64, by + 44, 1460, 100, "BEYOND CREATIVE WORK", HEAD, 44, True, INK, spacing=1.0)
label(sl, bx + 66, by + 142, 1540, 70,
      "My creative experience also grew through student organizations, academic projects, and collaborative work.",
      BODY, 16, False, MUTE, spacing=1.5)
blocks = [("Himpunan Mahasiswa PPP", "Head of Media & Branding", MSO_SHAPE.STAR_5_POINT),
          ("Pekan Seni Budaya IPB", "Vice Head / Staff DKV", MSO_SHAPE.STAR_4_POINT),
          ("Academic & Research Projects", "Visual communication · presentation · documentation", MSO_SHAPE.DIAMOND)]
for i, (title, sub, icon) in enumerate(blocks):
    cx, cy, cwd = bx + 66 + i * 560, by + 268, 520
    card = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cwd, 472)
    radius(card, 0.06); solid(card, AERO["soft"]); stroke(card, AERO["accent"], 1.75)
    top = shape(sl, MSO_SHAPE.RECTANGLE, cx, cy, cwd, 10)
    grad(top, AERO["bar_a"], AERO["pop"], 0); noline(top)
    ic = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx + 44, cy + 50, 88, 88)
    radius(ic, 0.22); grad(ic, AERO["bar_a"], AERO["bar_b"], 90); noline(ic)
    deco(sl, icon, cx + 66, cy + 72, 44, 44, "FFFFFF")
    label(sl, cx + 44, cy + 156, cwd - 88, 160, title, HEAD, 20, True, INK, spacing=1.2)
    label(sl, cx + 46, cy + 312, cwd - 92, 130, sub, BODY, 14, False, MUTE, spacing=1.4)
taskbar(sl, AERO, 8)

# --- 9 · skills & tools ---------------------------------------------------
sl = new(AERO)
bx, by, bw, bh = window(sl, 96, 108, 1728, 840, "skills_and_tools  —  Settings", AERO)
label(sl, bx + 64, by + 36, 1200, 100, "SKILLS & TOOLS", HEAD, 44, True, INK, spacing=1.0)
cols = [("Creative", ["Graphic Design", "Copywriting", "Content Creation", "Visual Communication"]),
        ("Marketing", ["Social Media", "Campaign Development", "Event Management", "Partnership"])]
for i, (title, items) in enumerate(cols):
    cx, cy = bx + 66 + i * 840, by + 142
    panel = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, 800, 356)
    radius(panel, 0.07); solid(panel, AERO["soft"]); stroke(panel, AERO["accent"], 1.75)
    label(sl, cx + 40, cy + 26, 700, 50, title, HEAD, 25, True, AERO["edge"])
    for j, it in enumerate(items):
        iy = cy + 92 + j * 62
        deco(sl, MSO_SHAPE.DIAMOND, cx + 40, iy + 10, 22, 22, AERO["bar_b"])
        label(sl, cx + 80, iy, 420, 42, it, BODY, 16, False, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        t = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx + 504, iy + 16, 252, 10)
        radius(t, 0.5); solid(t, "FFFFFF"); stroke(t, AERO["accent"], 1)
        f = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx + 504, iy + 16, 252 - j * 16, 10)
        radius(f, 0.5); grad(f, AERO["bar_a"], AERO["pop"], 0); noline(f)
tp = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx + 66, by + 534, 1574, 210)
radius(tp, 0.12); solid(tp, "FFFFFF"); stroke(tp, AERO["accent"], 1.75)
label(sl, bx + 106, by + 562, 500, 50, "Tools", HEAD, 25, True, AERO["edge"])
tools = ["Canva", "Adobe Illustrator", "Adobe Photoshop", "CapCut",
         "Microsoft Office", "Google Workspace", "Minitab"]
tx2 = bx + 106
for k, t in enumerate(tools):
    if k == 4:
        tx2 = bx + 106
    w = 48 + len(t) * 15 * 1.12
    ty = by + 626 if k < 4 else by + 680
    c = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, tx2, ty, w, 44)
    radius(c, 0.4); solid(c, AERO["chip"]); stroke(c, AERO["accent"], 1.25)
    label(sl, tx2, ty, w, 44, t, UI, 15, True, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tx2 += w + 14
taskbar(sl, AERO, 9)

# --- 10 · contact ---------------------------------------------------------
sl = new(PINK)
checker(sl, 0, 988, 1920)
for sx, sy, k, col, rot in [(150, 210, MSO_SHAPE.HEART, PINK["bar_a"], -10),
                            (1696, 246, MSO_SHAPE.STAR_5_POINT, PINK["accent"], 12),
                            (206, 772, MSO_SHAPE.SUN, PINK["pop"], 0),
                            (1736, 754, MSO_SHAPE.STAR_4_POINT, PINK["bar_a"], 0)]:
    sticker(sl, sx, sy, 100, k, col, rot)
bx, by, bw, bh = window(sl, 330, 112, 1260, 860, "new message!", PINK)
deco(sl, MSO_SHAPE.HEART, bx + 78, by + 36, 30, 28, PINK["bar_a"])
label(sl, bx + 120, by + 30, 600, 40, "1 new message", UI, 18, True, MUTE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
label(sl, bx + 76, by + 72, 1108, 140, "LET'S WORK", HEAD, 58, True, INK, spacing=1.0)
label(sl, bx + 76, by + 200, 1108, 140, "TOGETHER.", HEAD, 58, True, PINK["bar_a"], spacing=1.0)
rule(sl, bx + 80, by + 352, 360, PINK["pop"])
label(sl, bx + 78, by + 384, 1100, 190,
      "Thank you for taking the time to explore my work. I'm always open to new creative opportunities, collaborations, and projects.",
      BODY, 16, False, INK, spacing=1.55)
holo(sl, bx + 76, by + 540, 1108, 34, "thanks 4 stopping by  ·  let's create something")
label(sl, bx + 78, by + 610, 700, 60, "ZAHRA LEVINA", HEAD, 26, True, INK)
for i, (t, sub) in enumerate([("WhatsApp", "+62 8xx-xxxx-xxxx"), ("Email", "hello@email.com"),
                              ("LinkedIn", "linkedin.com/in/username")]):
    cx, cy = bx + 76 + i * 366, by + 672
    b = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, 340, 96)
    radius(b, 0.2); grad(b, PINK["bar_a"], PINK["bar_b"], 90); noline(b)
    label(sl, cx, cy + 18, 340, 34, t, UI, 17, True, "FFFFFF", PP_ALIGN.CENTER)
    label(sl, cx, cy + 52, 340, 30, sub, MONO, 11, False, "FFFFFF", PP_ALIGN.CENTER)
taskbar(sl, PINK, 10)

out = "/home/user/grill-me-skill/portfolio/zahra-levina-portfolio-2026.pptx"
prs.save(out)
print("saved", out, "| slides:", len(prs.slides._sldIdLst))
